# Проверенные рецепты для скилла pptx-restyle-universal. Скопируй в work/lib.py и импортируй из своих скриптов.
import copy, re, shutil, subprocess, tempfile
from pathlib import Path
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from lxml import etree

A = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
P = '{http://schemas.openxmlformats.org/presentationml/2006/main}'
R = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'
SHAPES = {P + 'sp', P + 'pic', P + 'grpSp', P + 'graphicFrame', P + 'cxnSp'}


def ptext(p):  # текст абзаца, <a:br> -> \n
    return ''.join(e.text or '' if e.tag == A + 't' else '\n' if e.tag == A + 'br' else '' for e in p.iter())


def xfrm(el):
    """(x, y, cx, cy, chx, chy, chcx, chcy) из spPr/grpSpPr/graphicFrame; None если позиция наследуется."""
    pr = el.find(P + 'spPr')
    if pr is None:
        pr = el.find(P + 'grpSpPr')
    xf = pr.find(A + 'xfrm') if pr is not None else el.find(P + 'xfrm')
    if xf is None or xf.find(A + 'off') is None:
        return None
    o, e = xf.find(A + 'off'), xf.find(A + 'ext')
    v = [int(o.get('x')), int(o.get('y')), int(e.get('cx')), int(e.get('cy'))]
    co, ce = xf.find(A + 'chOff'), xf.find(A + 'chExt')
    return v + ([int(co.get('x')), int(co.get('y')), int(ce.get('cx')), int(ce.get('cy'))] if co is not None else v)


def walk(tree, tf=(0, 0, 1.0, 1.0), parent=None):
    """Все фигуры (включая вложенные в группы) с АБСОЛЮТНЫМ bbox в EMU — координаты детей группы пересчитаны."""
    for el in tree:
        if el.tag not in SHAPES:
            continue
        nv = el.find('.//' + P + 'cNvPr')
        x = xfrm(el)
        ox, oy, sx, sy = tf
        bbox = (ox + x[0] * sx, oy + x[1] * sy, x[2] * sx, x[3] * sy) if x else None
        yield {'el': el, 'id': nv.get('id'), 'name': nv.get('name'), 'tag': el.tag.split('}')[1], 'bbox': bbox, 'parent': parent}
        if el.tag == P + 'grpSp' and x:
            ksx, ksy = (x[2] / x[6] if x[6] else 1), (x[3] / x[7] if x[7] else 1)
            yield from walk(el, (ox + x[0] * sx - x[4] * ksx * sx, oy + x[1] * sy - x[5] * ksy * sy, ksx * sx, ksy * sy), nv.get('id'))


def by_id(slide, sid):
    for s in walk(slide._element.cSld.spTree):
        if s['id'] == str(sid):
            return s['el']


def clone_slide(prs, src):
    """Копия слайда src в конец prs. ВАЖНО: не заменять cSld/spTree целиком — python-pptx кэширует
    slide.shapes на старый spTree, и всё, что потом добавишь через shapes.add_*, уйдёт в никуда."""
    new = prs.slides.add_slide(src.slide_layout)
    sp_new = new._element.cSld.spTree
    for ch in list(sp_new):
        sp_new.remove(ch)
    src_cSld = copy.deepcopy(src._element.cSld)
    for ch in list(src_cSld.spTree):
        sp_new.append(ch)
    bg = src_cSld.find(P + 'bg')
    if bg is not None:
        new._element.cSld.insert(0, bg)
    rmap = {}
    for rel in list(src.part.rels.values()):
        if rel.reltype in (RT.SLIDE_LAYOUT, RT.NOTES_SLIDE):
            continue
        rmap[rel.rId] = (new.part.relate_to(rel.target_ref, rel.reltype, is_external=True) if rel.is_external
                         else new.part.relate_to(rel.target_part, rel.reltype))
    for el in new._element.iter():
        for k, v in list(el.attrib.items()):
            if k.startswith(R) and v in rmap:
                el.set(k, rmap[v])
    return new


def delete_slide(prs, slide):
    lst = prs.slides._sldIdLst
    for sid in list(lst):
        if prs.part.related_part(sid.rId) is slide.part:
            rid = sid.rId
            lst.remove(sid)
            prs.part.drop_rel(rid)
            return


def set_paragraphs(el, items):
    """Заменить текст фигуры, СОХРАНИВ форматирование шаблона. items: список строк (абзацев).
    Абзац i получает стиль i-го НЕПУСТОГО абзаца шаблона (последний — для остальных); пустые абзацы-отбивки
    между стилями шаблона сохраняются; завершающий <a:br> шаблонного абзаца сохраняется; '<br>' внутри строки — перенос."""
    tb = el.find(P + 'txBody')
    tps = tb.findall(A + 'p')
    full = [i for i, p in enumerate(tps) if ptext(p).strip()] or [0]
    out = []
    for j, text in enumerate(items):
        k = full[min(j, len(full) - 1)]
        if 0 < j < len(full):
            out += [copy.deepcopy(tps[i]) for i in range(full[j - 1] + 1, full[j])]
        tp = tps[k]
        rpr = next((r.find(A + 'rPr') for r in tp.iter(A + 'r') if r.find(A + 'rPr') is not None), None)
        if rpr is None:
            rpr = tp.find(A + 'endParaRPr')
        p = copy.deepcopy(tp)
        end = p.find(A + 'endParaRPr')
        for ch in list(p):
            if ch.tag != A + 'pPr':
                p.remove(ch)
        for li, line in enumerate(re.split(r'<br\s*/?>', text)):
            if li:
                p.append(etree.Element(A + 'br'))
            r = etree.SubElement(p, A + 'r')
            if rpr is not None:
                rp = copy.deepcopy(rpr)
                rp.tag = A + 'rPr'
                r.append(rp)
            etree.SubElement(r, A + 't').text = line
        if ptext(tp).endswith('\n'):
            p.append(etree.Element(A + 'br'))
        if end is not None:
            p.append(end)
        out.append(p)
    if len(full) > 1:  # многоабзацный слот: якорь вверх, чтобы короткий текст не «уезжал» на иконку
        bp = tb.find(A + 'bodyPr')
        if bp is not None and bp.get('anchor') in ('ctr', 'b'):
            bp.set('anchor', 't')
    for p in tps:
        tb.remove(p)
    for p in out:
        tb.append(p)


def render(pptx, outdir, prefix, dpi=80):
    """pptx -> outdir/prefix.pdf + prefix-01.png …  (отдельный профиль LibreOffice, чтобы не ловить lock)."""
    outdir = Path(outdir)
    outdir.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as td:
        shutil.copy(pptx, f'{td}/in.pptx')
        subprocess.run(['soffice', f'-env:UserInstallation=file://{td}/lo', '--headless', '--convert-to', 'pdf',
                        '--outdir', td, f'{td}/in.pptx'], capture_output=True, timeout=600)
        shutil.copy(f'{td}/in.pdf', outdir / f'{prefix}.pdf')
    for f in outdir.glob(f'{prefix}-*.png'):
        f.unlink()
    subprocess.run(['pdftoppm', '-r', str(dpi), '-png', str(outdir / f'{prefix}.pdf'), str(outdir / f'{prefix}-t')], timeout=600)
    pngs = sorted(outdir.glob(f'{prefix}-t-*.png'), key=lambda p: int(re.findall(r'(\d+)\.png$', p.name)[0]))
    res = []
    for i, p in enumerate(pngs, 1):
        dst = outdir / f'{prefix}-{i:02d}.png'
        p.rename(dst)
        res.append(dst)
    return res


def pair(left, right, out, lt='TEMPLATE', rt='RESULT', w=760):
    from PIL import Image, ImageDraw
    a, b = Image.open(left).convert('RGB'), Image.open(right).convert('RGB')
    a, b = a.resize((w, int(a.height * w / a.width))), b.resize((w, int(b.height * w / b.width)))
    im = Image.new('RGB', (2 * w + 30, max(a.height, b.height) + 28), (235, 235, 235))
    d = ImageDraw.Draw(im)
    d.text((10, 6), lt, fill='black'); d.text((w + 20, 6), rt, fill='black')
    im.paste(a, (10, 22)); im.paste(b, (w + 20, 22))
    im.save(out, quality=85)


def words_outside(pdf, page_no, slide_w, bbox, text, tol=0.015):
    """Сколько слов текста фигуры отрендерилось вне её bbox (переполнение). page_no с 1."""
    out = subprocess.run(['pdftotext', '-f', str(page_no), '-l', str(page_no), '-bbox', str(pdf), '-'],
                         capture_output=True, text=True).stdout
    pw = float(re.search(r'<page width="([\d.]+)" height="([\d.]+)"', out).group(1))
    ph = float(re.search(r'<page width="([\d.]+)" height="([\d.]+)"', out).group(2))
    k = pw / slide_w
    x0, y0, x1, y1 = bbox[0] * k, bbox[1] * k, (bbox[0] + bbox[2]) * k, (bbox[1] + bbox[3]) * k
    pos = {}
    for m in re.finditer(r'<word xMin="([\d.]+)" yMin="([\d.]+)" xMax="([\d.]+)" yMax="([\d.]+)">(.*?)</word>', out):
        pos.setdefault(m.group(5).lower().strip('.,:;!?«»"()'), []).append(
            ((float(m.group(1)) + float(m.group(3))) / 2, (float(m.group(2)) + float(m.group(4))) / 2))
    from collections import Counter
    need = Counter(w.lower().strip('.,:;!?«»"()') for w in text.split() if len(w) >= 4)
    t = tol * ph
    missing = total = 0
    for w, n in need.items():
        if w not in pos:
            continue
        total += n
        inside = sum(1 for cx, cy in pos[w] if x0 - t <= cx <= x1 + t and y0 - t <= cy <= y1 + t)
        missing += max(0, n - inside)
    return missing, total   # переполнение, если missing / total > 0.1
